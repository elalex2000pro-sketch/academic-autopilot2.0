"""Servicios para IA Studio: extraccion de texto, generacion de PDF y resumen/infografia."""
from __future__ import annotations

import json
import os
import re
import uuid
from pathlib import Path
from typing import Any

import httpx
from pypdf import PdfReader
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.core.logging import get_logger
from app.db.base import utcnow
from app.models.ai_output import AiOutput
from app.models.document import Document
from app.models.user import User
from app.services.pdf_service import build_infographic_pdf, build_mindmap_pdf, build_summary_pdf

try:
    from docx import Document as DocxDocument
except Exception:  # pragma: no cover
    DocxDocument = None

try:
    from openpyxl import load_workbook
except Exception:  # pragma: no cover
    load_workbook = None

logger = get_logger(__name__)


def extract_text(document: Document) -> str:
    abs_path = os.path.join(settings.files_dir, document.storage_path)
    if not os.path.exists(abs_path):
        raise ValueError("Archivo no encontrado en almacenamiento")

    ext = Path(document.filename).suffix.lower()

    if ext == ".pdf":
        try:
            reader = PdfReader(abs_path)
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as exc:
            raise ValueError("No se pudo extraer texto del PDF") from exc
        return _sanitize(text)

    if ext in (".txt", ".md", ".csv"):
        try:
            with open(abs_path, encoding="utf-8", errors="replace") as f:
                return _sanitize(f.read())
        except Exception as exc:
            raise ValueError("No se pudo leer el archivo de texto") from exc

    if ext in (".docx",) and DocxDocument is not None:
        try:
            doc = DocxDocument(abs_path)
            text = "\n".join(p.text for p in doc.paragraphs)
        except Exception as exc:
            raise ValueError("No se pudo leer el archivo DOCX") from exc
        return _sanitize(text)

    if ext in (".xlsx", ".xlsm") and load_workbook is not None:
        try:
            workbook = load_workbook(abs_path, read_only=True, data_only=True)
            parts: list[str] = []
            for sheet in workbook.worksheets:
                parts.append(f"# Hoja: {sheet.title}")
                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index > 1500:
                        parts.append("... (hoja truncada)")
                        break
                    cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            workbook.close()
        except Exception as exc:
            raise ValueError("No se pudo leer el archivo de Excel") from exc
        return _sanitize("\n".join(parts))

    raise ValueError(f"Formato no soportado para resumen/infografia: {ext or 'desconocido'}")


def _sanitize(text: str) -> str:
    if not text:
        return ""
    # Limitar longitud para evitar prompts excesivos
    if len(text) > settings.ai_studio_max_chars:
        text = text[: settings.ai_studio_max_chars] + "\n\n[...texto truncado...]"
    return text.strip()


async def call_llm(user_text: str, kind: str) -> dict[str, Any]:
    """Llama a Ollama para obtener resumen o infografia estructurados."""
    if not settings.ai_configured:
        raise RuntimeError("IA no configurada")

    system_summary = (
        "Eres un asistente academico especializado en resumir documentos en espanol. "
        "Devuelve SIEMPRE un JSON valido con este esquema: {\"title\": \"...\", \"summary\": \"texto claro y conciso\", \"key_points\": [\"...\", \"...\"], \"keywords\": [\"...\"]}"
    )
    system_infog = (
        "Eres un asistente academico especializado en crear infografias. "
        "Devuelve SIEMPRE un JSON valido con este esquema: {\"title\": \"...\", \"subtitle\": \"...\", \"sections\": [{\"title\": \"...\", \"items\": [\"...\", \"...\"]}], \"stats\": [{\"label\": \"...\", \"value\": \"...\"]}, \"keywords\": [\"...\"]}"
    )
    system_mindmap = (
        "Eres un asistente academico especializado en mapas mentales. "
        "Devuelve SIEMPRE un JSON valido con este esquema: {\"title\": \"...\", \"center\": \"idea central en pocas palabras\", \"branches\": [{\"title\": \"rama principal\", \"children\": [\"detalle corto\", \"detalle corto\"]}]}. "
        "Incluye entre 4 y 8 ramas y entre 2 y 5 detalles por rama, todos cortos."
    )

    prompt = (
        "Este es el contenido del documento (posiblemente truncado):\n\n"
        f"{user_text}\n\n"
        "Genera un resultado academico, accionable y directo en espanol."
    )

    payload = {
        "model": settings.ai_model,
        "messages": [
            {
                "role": "system",
                "content": (
                    system_mindmap
                    if kind == "mindmap"
                    else (system_infog if kind == "infographic" else system_summary)
                ),
            },
            {"role": "user", "content": prompt},
        ],
        "temperature": settings.ai_temperature,
        "stream": False,
    }
    headers = {"Content-Type": "application/json"}
    if settings.ai_api_key:
        headers["Authorization"] = f"Bearer {settings.ai_api_key}"

    url = settings.ai_base_url.rstrip("/") + "/chat/completions"
    async with httpx.AsyncClient(timeout=settings.ai_timeout_seconds) as client:
        resp = await client.post(url, json=payload, headers=headers)
        resp.raise_for_status()
        data = resp.json()

    content = ((data.get("choices") or [{}])[0].get("message", {}) or {}).get("content") or ""
    content = content.strip()
    if not content:
        raise RuntimeError("Respuesta vacia del modelo")

    try:
        return _extract_json(content)
    except Exception as exc:
        logger.warning("JSON invalido del LLM: %s", exc)
        raise RuntimeError("El LLM no devolvio JSON valido") from exc


def _extract_json(content: str) -> dict[str, Any]:
    """Extrae el primer objeto JSON del texto, aunque el modelo anada prosa o ```."""
    text = content.strip()
    if text.startswith("```"):
        text = re.sub(r"^```[a-zA-Z]*\n?", "", text)
        text = re.sub(r"\n?```$", "", text).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    start = text.find("{")
    end = text.rfind("}")
    if start != -1 and end > start:
        return json.loads(text[start : end + 1])
    raise ValueError("No se encontro JSON en la respuesta")


def _heuristic_summary(doc_title: str, text: str) -> dict[str, Any]:
    sentences = [s.strip() for s in text.split("\n") if s.strip()]
    first = sentences[:5]
    return {
        "title": f"Resumen - {doc_title}",
        "summary": " ".join(first) or doc_title,
        "key_points": first,
        "keywords": [],
    }


def _heuristic_infographic(doc_title: str, text: str) -> dict[str, Any]:
    first = text.split("\n")[:6]
    return {
        "title": f"Infografia - {doc_title}",
        "subtitle": "Generada por motor de reglas",
        "sections": [{"title": "Ideas principales", "items": [x.strip() for x in first if x.strip()][:8]}],
        "stats": [],
        "keywords": [],
    }


def _heuristic_mindmap(doc_title: str, text: str) -> dict[str, Any]:
    lines = [x.strip() for x in text.split("\n") if x.strip()][:6]
    return {
        "title": f"Mapa mental - {doc_title}",
        "center": doc_title[:80],
        "branches": [{"title": line[:60], "children": []} for line in lines],
    }


async def generate(
    db: AsyncSession,
    user: User,
    document: Document,
    kind: str,
) -> AiOutput:
    source = "ai"
    text = extract_text(document)
    try:
        data = await call_llm(text, kind)
    except Exception as exc:  # noqa: BLE001
        logger.warning("AI Studio usando heuristico: %s", exc)
        source = "heuristic"
        if kind == "summary":
            data = _heuristic_summary(document.filename, text)
        elif kind == "mindmap":
            data = _heuristic_mindmap(document.filename, text)
        else:
            data = _heuristic_infographic(document.filename, text)

    if kind == "summary":
        pdf_bytes = build_summary_pdf(
            data.get("title", "Resumen"),
            {"filename": document.filename, "created_at": utcnow().strftime("%Y-%m-%d %H:%M")},
            data.get("summary", ""),
            data.get("key_points", []) or [],
        )
    elif kind == "mindmap":
        pdf_bytes = build_mindmap_pdf(
            data.get("title", "Mapa mental"),
            data.get("center", ""),
            data.get("branches", []) or [],
        )
    else:
        pdf_bytes = build_infographic_pdf(
            data.get("title", "Infografia"),
            data.get("subtitle", ""),
            data.get("sections", []) or [],
            data.get("stats", []) or [],
        )

    out_doc = await _save_generated_pdf(db, user, document, kind, data, pdf_bytes)
    ai_out = AiOutput(
        user_id=user.id,
        source_document_id=document.id,
        output_document_id=out_doc.id,
        kind=kind,
        title=str(data.get("title", f"{kind} - {document.filename}"))[:500],
        content=(data.get("summary") if kind == "summary" else json.dumps(data))[:200000],
        source=source,
    )
    db.add(ai_out)
    await db.commit()
    await db.refresh(ai_out)
    return ai_out


def _write_bytes(path: str, data: bytes) -> None:
    with open(path, "wb") as f:
        f.write(data)


async def _save_generated_pdf(
    db: AsyncSession,
    user: User,
    document: Document,
    kind: str,
    data: dict[str, Any],
    pdf_bytes: bytes,
) -> Document:
    os.makedirs(settings.files_dir, exist_ok=True)
    base_dir = os.path.join(settings.files_dir, str(user.id), "generated")
    os.makedirs(base_dir, exist_ok=True)
    fname = f"{kind}-{uuid.uuid4().hex}.pdf"
    abs_path = os.path.join(base_dir, fname)
    _write_bytes(abs_path, pdf_bytes)
    rel_path = os.path.join(str(user.id), "generated", fname)

    out_doc = Document(
        id=uuid.uuid4(),
        user_id=user.id,
        course_id=document.course_id,
        filename=f"{data.get('title','Documento')} [{kind}].pdf",
        content_type="application/pdf",
        size_bytes=len(pdf_bytes),
        storage_path=rel_path,
    )
    db.add(out_doc)
    await db.commit()
    await db.refresh(out_doc)
    return out_doc
