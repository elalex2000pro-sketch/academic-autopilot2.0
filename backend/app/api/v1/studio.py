"""Endpoints del IA Studio: generar resumenes, infografias y mapas mentales (PDF)."""
from __future__ import annotations

import io
import json
import os
import uuid
from urllib.parse import quote

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, HTTPException, Response, status
from sqlalchemy import delete, select

from app.api.deps import CurrentUser, DbSession
from app.core.config import settings
from app.core.logging import get_logger
from app.models.ai_output import AiOutput
from app.models.document import Document
from app.schemas.studio import AiOutputRead, StudioGenerateRequest
from app.services import ai_studio_service
from app.services.realtime import notify

logger = get_logger(__name__)

router = APIRouter(prefix="/studio", tags=["studio"])


@router.post("/generate", response_model=AiOutputRead, status_code=status.HTTP_201_CREATED)
async def generate(
    payload: StudioGenerateRequest, current_user: CurrentUser, db: DbSession
) -> AiOutput:
    document = await _get_owned_document(db, current_user.id, payload.document_id)
    try:
        output = await ai_studio_service.generate(db, current_user, document, payload.kind)
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    except Exception as exc:  # noqa: BLE001
        logger.error("Fallo al generar con IA: %s", exc)
        raise HTTPException(
            status_code=status.HTTP_502_BAD_GATEWAY,
            detail="No se pudo generar la salida de IA",
        ) from exc
    notify(current_user.id, "studio", {"action": "created", "id": str(output.id)})
    return output


@router.get("/outputs", response_model=list[AiOutputRead])
async def list_outputs(
    current_user: CurrentUser,
    db: DbSession,
    kind: str | None = None,
) -> list[AiOutput]:
    stmt = select(AiOutput).where(AiOutput.user_id == current_user.id)
    if kind is not None:
        stmt = stmt.where(AiOutput.kind == kind)
    stmt = stmt.order_by(AiOutput.created_at.desc())
    result = await db.execute(stmt)
    return list(result.scalars().all())


@router.delete("/outputs/{output_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
async def delete_output(
    output_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> None:
    result = await db.execute(
        select(AiOutput).where(AiOutput.id == output_id, AiOutput.user_id == current_user.id)
    )
    output = result.scalar_one_or_none()
    if output is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Salida no encontrada")

    if output.output_document_id is not None:
        doc_result = await db.execute(
            select(Document).where(
                Document.id == output.output_document_id,
                Document.user_id == current_user.id,
            )
        )
        doc = doc_result.scalar_one_or_none()
        if doc is not None:
            abs_path = os.path.join(settings.files_dir, doc.storage_path)
            await db.delete(doc)
            await db.execute(
                delete(AiOutput).where(AiOutput.id == output_id)
            )
            await db.commit()
            _safe_remove(abs_path)
            notify(current_user.id, "studio", {"action": "deleted", "id": str(output_id)})
            return

    await db.delete(output)
    await db.commit()
    notify(current_user.id, "studio", {"action": "deleted", "id": str(output_id)})


@router.get("/outputs/{output_id}/word")
async def export_output_word(
    output_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> Response:
    result = await db.execute(
        select(AiOutput).where(AiOutput.id == output_id, AiOutput.user_id == current_user.id)
    )
    output = result.scalar_one_or_none()
    if output is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Salida no encontrada")

    kind_label = {"summary": "Resumen", "infographic": "Infografia", "mindmap": "Mapa mental"}.get(
        output.kind, output.kind
    )
    document = DocxDocument()
    title = document.add_heading(output.title or kind_label, level=0)
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x31, 0x2E, 0x81)

    meta = document.add_paragraph()
    run = meta.add_run(
        f"Academic Autopilot · {kind_label} · {output.created_at.strftime('%d/%m/%Y %H:%M')}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    payload = _parse_json(output.content)
    if output.kind == "summary" or payload is None:
        for block in (output.content or "").split("\n"):
            if block.strip():
                document.add_paragraph(block.strip())
    elif output.kind == "mindmap":
        center = payload.get("center") or payload.get("title") or ""
        if center:
            paragraph = document.add_paragraph()
            run = paragraph.add_run(str(center))
            run.bold = True
        for branch in payload.get("branches", []) or []:
            document.add_heading(str(branch.get("title", "")), level=2)
            for child in branch.get("children", []) or []:
                document.add_paragraph(str(child), style="List Bullet")
    else:
        subtitle = payload.get("subtitle")
        if subtitle:
            document.add_paragraph(str(subtitle))
        for section in payload.get("sections", []) or []:
            document.add_heading(str(section.get("title", "")), level=2)
            for item in section.get("items", []) or []:
                document.add_paragraph(str(item), style="List Bullet")

    buffer = io.BytesIO()
    document.save(buffer)
    encoded = quote(f"{output.title or kind_label}.docx")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename=\"salida.docx\"; filename*=UTF-8''{encoded}"
        },
    )


def _parse_json(content: str) -> dict | None:
    try:
        data = json.loads(content)
    except Exception:  # noqa: BLE001
        return None
    return data if isinstance(data, dict) else None


async def _get_owned_document(
    db: DbSession, user_id: uuid.UUID, document_id: uuid.UUID
) -> Document:
    result = await db.execute(
        select(Document).where(Document.id == document_id, Document.user_id == user_id)
    )
    document = result.scalar_one_or_none()
    if document is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado")
    return document


def _safe_remove(path: str) -> None:
    try:
        if os.path.exists(path):
            os.remove(path)
    except OSError as exc:  # pragma: no cover
        logger.warning("No se pudo borrar %s: %s", path, exc)
