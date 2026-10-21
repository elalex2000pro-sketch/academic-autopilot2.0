"""Exportaciones a Excel (tareas y cuadernos)."""
from __future__ import annotations

import io
from datetime import datetime

from fastapi import APIRouter, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import selectinload

from app.api.deps import CurrentUser, DbSession
from app.models.notebook import Notebook
from app.models.task import Task

router = APIRouter(prefix="/exports", tags=["exports"])

_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _xlsx_response(workbook: Workbook, filename: str) -> Response:
    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def _style_header(sheet, columns: list[str], widths: list[int]) -> None:
    for index, (name, width) in enumerate(zip(columns, widths), start=1):
        cell = sheet.cell(row=1, column=index, value=name)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"


def _fmt(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


_STATUS = {
    "pending": "Pendiente",
    "in_progress": "En curso",
    "completed": "Completada",
    "archived": "Archivada",
}
_PRIORITY = {
    "urgent": "Urgente",
    "long_project": "Proyecto largo",
    "short_task": "Tarea corta",
}


@router.get("/tasks.xlsx")
async def export_tasks_excel(current_user: CurrentUser, db: DbSession) -> Response:
    result = await db.execute(
        select(Task)
        .options(selectinload(Task.course))
        .where(Task.user_id == current_user.id)
        .order_by(Task.due_at.asc().nullslast(), Task.created_at.desc())
    )
    tasks = list(result.scalars().all())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Tareas"
    _style_header(
        sheet,
        [
            "Tarea",
            "Asignatura",
            "Estado",
            "Prioridad",
            "Entrega",
            "Minutos estimados",
            "Descripcion",
        ],
        [44, 26, 14, 16, 18, 18, 60],
    )
    for row, task in enumerate(tasks, start=2):
        sheet.cell(row=row, column=1, value=task.title)
        sheet.cell(row=row, column=2, value=task.course.name if task.course else "Sin asignatura")
        sheet.cell(row=row, column=3, value=_STATUS.get(task.status.value, task.status.value))
        sheet.cell(
            row=row, column=4, value=_PRIORITY.get(task.priority.value, task.priority.value)
        )
        sheet.cell(row=row, column=5, value=_fmt(task.due_at))
        sheet.cell(row=row, column=6, value=task.estimated_minutes)
        sheet.cell(row=row, column=7, value=task.description or "")
    return _xlsx_response(workbook, "tareas-academic-autopilot.xlsx")


@router.get("/notebooks.xlsx")
async def export_notebooks_excel(current_user: CurrentUser, db: DbSession) -> Response:
    result = await db.execute(
        select(Notebook)
        .where(Notebook.user_id == current_user.id)
        .order_by(Notebook.updated_at.desc())
    )
    notebooks = list(result.scalars().all())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Cuadernos"
    _style_header(
        sheet, ["Cuaderno", "Actualizado", "Contenido"], [36, 18, 100]
    )
    for row, notebook in enumerate(notebooks, start=2):
        sheet.cell(row=row, column=1, value=notebook.title)
        sheet.cell(row=row, column=2, value=_fmt(notebook.updated_at))
        cell = sheet.cell(row=row, column=3, value=notebook.content or "")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    return _xlsx_response(workbook, "cuadernos-academic-autopilot.xlsx")
