"""Endpoints de cuadernos de notas y su exportacion a Word."""
from __future__ import annotations

import io
import uuid
from typing import Literal
from urllib.parse import quote

from docx import Document as DocxDocument
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, HTTPException, Response, status
from pydantic import BaseModel
from sqlalchemy import select

from app.api.deps import CurrentUser, DbSession
from app.models.book import Book
from app.models.document import Document
from app.models.notebook import Notebook, NotebookSource
from app.models.shared_class import ClassMembership
from app.schemas.notebook import (
    NotebookCreate,
    NotebookRead,
    NotebookSourceCreate,
    NotebookSourceRead,
    NotebookUpdate,
)

router = APIRouter(prefix="/notebooks", tags=["notebooks"])


class NotebookShareUpdate(BaseModel):
    class_id: uuid.UUID
    permission: Literal["read", "edit"] = "read"

_INK = RGBColor(0x31, 0x2E, 0x81)
_GREY = RGBColor(0x64, 0x74, 0x8B)


def _docx_response(doc: DocxDocument, filename: str) -> Response:
    buffer = io.BytesIO()
    doc.save(buffer)
    ascii_name = "cuaderno.docx"
    encoded = quote(f"{filename}.docx")
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={
            "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{encoded}"
        },
    )


@router.get("", response_model=list[NotebookRead])
async def list_notebooks(current_user: CurrentUser, db: DbSession) -> list[Notebook]:
    result = await db.execute(
        select(Notebook)
        .where(Notebook.user_id == current_user.id)
        .order_by(Notebook.updated_at.desc())
    )
    return list(result.scalars().all())


@router.post("", response_model=NotebookRead, status_code=status.HTTP_201_CREATED)
async def create_notebook(
    payload: NotebookCreate, current_user: CurrentUser, db: DbSession
) -> Notebook:
    notebook = Notebook(
        user_id=current_user.id,
        title=payload.title,
        color=payload.color,
        content=payload.content,
    )
    db.add(notebook)
    await db.commit()
    await db.refresh(notebook)
    return notebook


@router.get("/{notebook_id}", response_model=NotebookRead)
async def get_notebook(
    notebook_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> Notebook:
    return await _get_accessible(notebook_id, current_user.id, db)


def _source_read(
    source: NotebookSource,
    book: Book | None,
    document: Document | None,
) -> NotebookSourceRead:
    if book is not None:
        subtitle_bits = [book.author] if book.author else []
        subtitle_bits.append(f"{book.total_sections} secciones")
        return NotebookSourceRead(
            id=source.id,
            notebook_id=source.notebook_id,
            kind="book",
            title=book.title,
            subtitle=" · ".join(subtitle_bits),
            book_id=book.id,
            format=book.format.upper(),
            size_bytes=book.size_bytes,
            created_at=source.created_at,
        )
    assert document is not None
    extension = document.filename.rsplit(".", 1)[-1].upper() if "." in document.filename else None
    return NotebookSourceRead(
        id=source.id,
        notebook_id=source.notebook_id,
        kind="file",
        title=document.filename,
        subtitle=document.content_type or "Archivo",
        document_id=document.id,
        format=extension,
        size_bytes=document.size_bytes,
        created_at=source.created_at,
    )


@router.get("/{notebook_id}/sources", response_model=list[NotebookSourceRead])
async def list_notebook_sources(
    notebook_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> list[NotebookSourceRead]:
    """Fuentes (libros o archivos) elegidas para este cuaderno."""
    await _get_accessible(notebook_id, current_user.id, db)
    result = await db.execute(
        select(NotebookSource)
        .where(NotebookSource.notebook_id == notebook_id)
        .order_by(NotebookSource.created_at.asc())
    )
    sources = list(result.scalars().all())
    book_ids = [source.book_id for source in sources if source.book_id]
    document_ids = [source.document_id for source in sources if source.document_id]
    books: dict[uuid.UUID, Book] = {}
    documents: dict[uuid.UUID, Document] = {}
    if book_ids:
        rows = await db.execute(select(Book).where(Book.id.in_(book_ids)))
        books = {book.id: book for book in rows.scalars().all()}
    if document_ids:
        rows = await db.execute(select(Document).where(Document.id.in_(document_ids)))
        documents = {document.id: document for document in rows.scalars().all()}
    return [
        _source_read(
            source,
            books.get(source.book_id) if source.book_id else None,
            documents.get(source.document_id) if source.document_id else None,
        )
        for source in sources
        if (source.book_id and source.book_id in books)
        or (source.document_id and source.document_id in documents)
    ]


@router.post(
    "/{notebook_id}/sources",
    response_model=NotebookSourceRead,
    status_code=status.HTTP_201_CREATED,
)
async def add_notebook_source(
    notebook_id: uuid.UUID,
    payload: NotebookSourceCreate,
    current_user: CurrentUser,
    db: DbSession,
) -> NotebookSourceRead:
    """Añade un libro o archivo ya subido como fuente del cuaderno."""
    await _get_owned(notebook_id, current_user.id, db)
    if (payload.book_id is None) == (payload.document_id is None):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Elige un libro o un archivo (solo uno).",
        )
    book: Book | None = None
    document: Document | None = None
    if payload.book_id:
        book = await db.get(Book, payload.book_id)
        if book is None or book.user_id != current_user.id:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Libro no encontrado"
            )
        duplicate = await db.execute(
            select(NotebookSource.id).where(
                NotebookSource.notebook_id == notebook_id,
                NotebookSource.book_id == book.id,
            )
        )
    else:
        document = await db.get(Document, payload.document_id)
        if document is None or document.user_id != current_user.id:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND, detail="Archivo no encontrado"
            )
        duplicate = await db.execute(
            select(NotebookSource.id).where(
                NotebookSource.notebook_id == notebook_id,
                NotebookSource.document_id == document.id,
            )
        )
    if duplicate.scalar_one_or_none() is not None:
        raise HTTPException(
            status_code=status.HTTP_409_CONFLICT,
            detail="Esa fuente ya está en el cuaderno.",
        )
    source = NotebookSource(
        notebook_id=notebook_id,
        book_id=book.id if book else None,
        document_id=document.id if document else None,
        added_by=current_user.id,
    )
    db.add(source)
    await db.commit()
    await db.refresh(source)
    return _source_read(source, book, document)


@router.delete(
    "/{notebook_id}/sources/{source_id}",
    status_code=status.HTTP_204_NO_CONTENT,
    response_model=None,
)
async def remove_notebook_source(
    notebook_id: uuid.UUID,
    source_id: uuid.UUID,
    current_user: CurrentUser,
    db: DbSession,
) -> None:
    """Quita una fuente del cuaderno sin borrar el libro ni el archivo."""
    await _get_owned(notebook_id, current_user.id, db)
    source = await db.get(NotebookSource, source_id)
    if source is None or source.notebook_id != notebook_id:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Fuente no encontrada"
        )
    await db.delete(source)
    await db.commit()


@router.patch("/{notebook_id}", response_model=NotebookRead)
async def update_notebook(
    notebook_id: uuid.UUID,
    payload: NotebookUpdate,
    current_user: CurrentUser,
    db: DbSession,
) -> Notebook:
    notebook = await _get_accessible(notebook_id, current_user.id, db, write=True)
    data = payload.model_dump(exclude_unset=True)
    for key, value in data.items():
        setattr(notebook, key, value)
    await db.commit()
    await db.refresh(notebook)
    return notebook


@router.put("/{notebook_id}/share", response_model=NotebookRead)
async def share_notebook(
    notebook_id: uuid.UUID,
    payload: NotebookShareUpdate,
    current_user: CurrentUser,
    db: DbSession,
) -> Notebook:
    notebook = await _get_owned(notebook_id, current_user.id, db)
    membership = await db.execute(
        select(ClassMembership).where(
            ClassMembership.class_id == payload.class_id,
            ClassMembership.user_id == current_user.id,
        )
    )
    if membership.scalar_one_or_none() is None:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Solo puedes compartir con una clase a la que perteneces",
        )
    notebook.shared_class_id = payload.class_id
    notebook.share_permission = payload.permission
    await db.commit()
    await db.refresh(notebook)
    return notebook


@router.delete("/{notebook_id}/share", response_model=NotebookRead)
async def unshare_notebook(
    notebook_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> Notebook:
    notebook = await _get_owned(notebook_id, current_user.id, db)
    notebook.shared_class_id = None
    notebook.share_permission = "read"
    await db.commit()
    await db.refresh(notebook)
    return notebook


@router.delete("/{notebook_id}", status_code=status.HTTP_204_NO_CONTENT, response_model=None)
async def delete_notebook(
    notebook_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> None:
    notebook = await _get_owned(notebook_id, current_user.id, db)
    await db.delete(notebook)
    await db.commit()


@router.get("/{notebook_id}/word")
async def export_notebook_word(
    notebook_id: uuid.UUID, current_user: CurrentUser, db: DbSession
) -> Response:
    notebook = await _get_owned(notebook_id, current_user.id, db)
    document = DocxDocument()

    title = document.add_heading(notebook.title or "Cuaderno", level=0)
    for run in title.runs:
        run.font.color.rgb = _INK

    meta = document.add_paragraph()
    run = meta.add_run(
        f"Academic Autopilot · Actualizado: {notebook.updated_at.strftime('%d/%m/%Y %H:%M')}"
    )
    run.font.size = Pt(9)
    run.font.color.rgb = _GREY

    for block in (notebook.content or "").split("\n"):
        line = block.rstrip()
        if not line:
            document.add_paragraph()
            continue
        if line.startswith("### "):
            document.add_heading(line[4:], level=3)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith(("- ", "* ")):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)

    return _docx_response(document, notebook.title or "cuaderno")


async def _get_owned(notebook_id: uuid.UUID, user_id: uuid.UUID, db: DbSession) -> Notebook:
    notebook = await db.get(Notebook, notebook_id)
    if notebook is None or notebook.user_id != user_id:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Cuaderno no encontrado"
        )
    return notebook


async def _get_accessible(
    notebook_id: uuid.UUID,
    user_id: uuid.UUID,
    db: DbSession,
    *,
    write: bool = False,
) -> Notebook:
    """Cuaderno propio o compartido con una clase a la que pertenece el usuario."""
    notebook = await db.get(Notebook, notebook_id)
    if notebook is None:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND, detail="Cuaderno no encontrado"
        )
    if notebook.user_id == user_id:
        return notebook
    if notebook.shared_class_id is not None:
        result = await db.execute(
            select(ClassMembership).where(
                ClassMembership.class_id == notebook.shared_class_id,
                ClassMembership.user_id == user_id,
            )
        )
        if result.scalar_one_or_none() is not None:
            if not write or notebook.share_permission == "edit":
                return notebook
    raise HTTPException(
        status_code=status.HTTP_404_NOT_FOUND, detail="Cuaderno no encontrado"
    )
